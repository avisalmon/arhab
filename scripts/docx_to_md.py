"""Convert the exported trip-notes .docx into a Markdown reference file.

Usage (from repo root, using the project's `env` virtualenv):
    .\\env\\Scripts\\python.exe scripts\\docx_to_md.py docs\\source-trip-notes.docx docs\\source-trip-notes.md

Walks the document body in order so paragraphs and tables come out
interleaved the way they appear in the original doc, rather than all
paragraphs then all tables (which is how python-docx exposes them by
default via .paragraphs / .tables).
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(parent):
    """Yield each paragraph and table in the document, in document order."""
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def paragraph_to_md(p: Paragraph) -> str:
    text = p.text.strip()
    if not text:
        return ""
    style = (p.style.name or "").lower() if p.style else ""
    if style.startswith("heading 1") or style == "title":
        return f"## {text}"
    if style.startswith("heading 2"):
        return f"### {text}"
    if style.startswith("heading 3"):
        return f"#### {text}"
    if style.startswith("list") or p._p.pPr is not None and p._p.pPr.numPr is not None:
        return f"- {text}"
    return text


def table_to_md(t: Table) -> str:
    rows = [[cell.text.strip().replace("\n", "<br>") for cell in row.cells] for row in t.rows]
    if not rows:
        return ""
    lines = []
    header, *body = rows
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def convert(src: Path, dst: Path):
    doc = Document(str(src))
    out_lines = [
        "<!--",
        "  Auto-generated from the trip notes Google Doc export.",
        "  Do not hand-edit — re-export the .docx and re-run",
        "  scripts/docx_to_md.py instead, or edits will be lost.",
        "-->",
        "",
        "# Trip Notes (source of truth)",
        "",
    ]
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            md = paragraph_to_md(block)
            if md:
                out_lines.append(md)
                out_lines.append("")
        elif isinstance(block, Table):
            md = table_to_md(block)
            if md:
                out_lines.append(md)
                out_lines.append("")
    dst.write_text("\n".join(out_lines), encoding="utf-8")
    print(f"Wrote {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python docx_to_md.py <input.docx> <output.md>")
        sys.exit(1)
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
